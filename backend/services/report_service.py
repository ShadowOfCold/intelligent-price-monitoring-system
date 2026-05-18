import os
import re
import tempfile
from datetime import datetime

import pandas as pd
import plotly.express as px
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from backend.models.correlation_analysis_result import CorrelationAnalysisResult
from backend.models.detected_anomaly import DetectedAnomaly
from backend.models.price import Price
from backend.models.price_forecast import PriceForecast
from backend.models.product import Product
from backend.models.report import Report
from backend.utils.datetime_utils import get_current_datetime

REPORTS_DIR = "data/reports"
FONT_PATH = "backend/assets/fonts/DejaVuSans.ttf"
FONT_NAME = "DejaVu"


def ensure_reports_directory():
    os.makedirs(REPORTS_DIR, exist_ok=True)


def register_pdf_font():
    if os.path.exists(FONT_PATH):
        try:
            pdfmetrics.registerFont(TTFont(FONT_NAME, FONT_PATH))
            return FONT_NAME
        except Exception:
            return "Helvetica"

    return "Helvetica"


def safe_filename(value: str) -> str:
    value = value.strip()
    value = re.sub(r'[\\/:*?"<>|]', "_", value)
    value = value.replace(" ", "_")
    return value


def normalize_numeric_series(series: pd.Series) -> pd.Series:
    return pd.to_numeric(
        series.astype(str)
        .str.replace(" ", "", regex=False)
        .str.replace("\u00a0", "", regex=False)
        .str.replace(",", ".", regex=False),
        errors="coerce"
    )


def autofit_excel_columns(worksheet):
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value).lower()

        for cell in column_cells:
            cell.alignment = Alignment(
                horizontal="right"
            )

            if isinstance(cell.value, (int, float)):
                if "id" in header:
                    cell.number_format = "0"
                else:
                    cell.number_format = "# ##0.00"

            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        worksheet.column_dimensions[column_letter].width = max(
            min(max_length + 2, 60),
            15
        )


def get_product(db: Session, product_id: int):
    return db.query(Product).filter(Product.id == product_id).first()


def get_prices_dataframe(db: Session, product_id: int):
    prices = (
        db.query(Price)
        .filter(Price.product_id == product_id)
        .order_by(Price.checked_at)
        .all()
    )

    if not prices:
        return pd.DataFrame()

    return pd.DataFrame(
        [
            {
                "ID": int(price.id),
                "Цена": float(price.price),
                "Дата": price.checked_at.strftime("%d.%m.%Y %H:%M"),
                "ID карточки": int(price.store_product_id) if price.store_product_id else None,
                "ID документа": int(price.document_id) if price.document_id else None,
            }
            for price in prices
        ]
    )


def get_anomalies_dataframe(db: Session, product_id: int):
    anomalies = (
        db.query(DetectedAnomaly, Price)
        .join(Price, DetectedAnomaly.price_id == Price.id)
        .filter(Price.product_id == product_id)
        .order_by(Price.checked_at)
        .all()
    )

    if not anomalies:
        return pd.DataFrame()

    anomaly_type_map = {
        "price_outlier": "Нетипичное значение цены",
        "z_score_outlier": "Статистический выброс",
        "price_spike": "Резкий рост цены",
        "price_drop": "Резкое падение цены",
        "price_fixation": "Длительная фиксация цены",
        "synchronized_change": "Синхронное изменение цены",
    }

    risk_level_map = {
        "low": "Низкий",
        "medium": "Средний",
        "high": "Высокий",
    }

    return pd.DataFrame(
        [
            {
                "ID": int(anomaly.id),
                "ID цены": int(anomaly.price_id),
                "Цена": float(price.price),
                "Дата цены": price.checked_at.strftime("%d.%m.%Y %H:%M"),
                "Тип аномалии": anomaly_type_map.get(
                    anomaly.anomaly_type,
                    anomaly.anomaly_type
                ),
                "Уровень риска": risk_level_map.get(
                    anomaly.risk_level,
                    anomaly.risk_level
                ),
                "Оценка": round(float(anomaly.anomaly_score), 4),
                "Дата выявления": anomaly.detected_at.strftime("%d.%m.%Y %H:%M"),
            }
            for anomaly, price in anomalies
        ]
    )


def get_correlations_dataframe(db: Session, product_id: int):
    correlations = (
        db.query(CorrelationAnalysisResult)
        .filter(CorrelationAnalysisResult.product_id == product_id)
        .order_by(CorrelationAnalysisResult.created_at.desc())
        .all()
    )

    if not correlations:
        return pd.DataFrame()

    risk_level_map = {
        "low": "Низкий",
        "medium": "Средний",
        "high": "Высокий",
    }

    return pd.DataFrame(
        [
            {
                "ID": int(correlation.id),
                "ID первой карточки": int(correlation.first_store_product_id),
                "ID второй карточки": int(correlation.second_store_product_id),
                "Коэффициент корреляции": round(float(correlation.correlation_value), 4),
                "Уровень риска": risk_level_map.get(
                    correlation.risk_level,
                    correlation.risk_level
                ),
                "Начало периода": correlation.period_start.strftime("%d.%m.%Y"),
                "Конец периода": correlation.period_end.strftime("%d.%m.%Y"),
                "Дата анализа": correlation.created_at.strftime("%d.%m.%Y %H:%M"),
            }
            for correlation in correlations
        ]
    )


def get_forecasts_dataframe(db: Session, product_id: int):
    forecasts = (
        db.query(PriceForecast)
        .filter(PriceForecast.product_id == product_id)
        .order_by(PriceForecast.forecast_date)
        .all()
    )

    if not forecasts:
        return pd.DataFrame()

    return pd.DataFrame(
        [
            {
                "ID": int(forecast.id),
                "ID карточки": int(forecast.store_product_id) if forecast.store_product_id else None,
                "Дата прогноза": forecast.forecast_date.strftime("%d.%m.%Y"),
                "Прогнозируемая цена": float(forecast.predicted_price),
                "Дата формирования": forecast.created_at.strftime("%d.%m.%Y %H:%M"),
            }
            for forecast in forecasts
        ]
    )


def get_period(prices_df: pd.DataFrame):
    if prices_df.empty:
        now = get_current_datetime()
        return now, now

    dates = pd.to_datetime(
        prices_df["Дата"],
        format="%d.%m.%Y %H:%M",
        errors="coerce"
    ).dropna()

    if dates.empty:
        now = get_current_datetime()
        return now, now

    return dates.min().to_pydatetime(), dates.max().to_pydatetime()


def write_dataframe_to_excel(writer, dataframe, sheet_name: str):
    if dataframe.empty:
        dataframe = pd.DataFrame([{"Сообщение": "Данные отсутствуют"}])

    dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    worksheet = writer.sheets[sheet_name]
    autofit_excel_columns(worksheet)


def add_price_chart_to_excel(writer, sheet_name: str):
    worksheet = writer.sheets[sheet_name]

    if worksheet.max_row < 3:
        return

    chart = LineChart()
    chart.title = "История цен"
    chart.y_axis.title = "Цена, руб."
    chart.x_axis.title = "Дата"

    chart.y_axis.numFmt = "# ##0"
    chart.y_axis.delete = False
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblPos = "nextTo"

    data = Reference(
        worksheet,
        min_col=2,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=3,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 12
    chart.width = 24

    worksheet.add_chart(chart, "G2")


def add_forecast_chart_to_excel(writer, sheet_name: str):
    worksheet = writer.sheets[sheet_name]

    if worksheet.max_row < 3:
        return

    chart = LineChart()
    chart.title = "Прогноз цен"
    chart.y_axis.title = "Цена, руб."
    chart.x_axis.title = "Дата"

    chart.y_axis.numFmt = "# ##0"
    chart.y_axis.delete = False
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblPos = "nextTo"

    data = Reference(
        worksheet,
        min_col=4,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=3,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 12
    chart.width = 24

    worksheet.add_chart(chart, "G2")


def create_report_record(
    db: Session,
    file_name: str,
    file_path: str,
    period_start: datetime,
    period_end: datetime
):
    report = Report(
        file_name=file_name,
        file_path=file_path,
        created_at=get_current_datetime(),
        period_start=period_start,
        period_end=period_end,
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    return report


def generate_excel_report(db: Session, product_id: int):
    ensure_reports_directory()

    product = get_product(db, product_id)

    if product is None:
        raise ValueError("Товар не найден")

    product_name = safe_filename(product.name)
    created_at_text = get_current_datetime().strftime("%d-%m-%Y_%H-%M-%S")

    file_name = f"Отчёт_{product_name}_{created_at_text}.xlsx"
    file_path = os.path.join(REPORTS_DIR, file_name)

    prices_df = get_prices_dataframe(db, product_id)
    anomalies_df = get_anomalies_dataframe(db, product_id)
    correlations_df = get_correlations_dataframe(db, product_id)
    forecasts_df = get_forecasts_dataframe(db, product_id)

    period_start, period_end = get_period(prices_df)

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        write_dataframe_to_excel(writer, prices_df, "История цен")
        write_dataframe_to_excel(writer, anomalies_df, "Аномалии")
        write_dataframe_to_excel(writer, correlations_df, "Корреляции")
        write_dataframe_to_excel(writer, forecasts_df, "Прогнозы")

        add_price_chart_to_excel(writer, "История цен")
        add_forecast_chart_to_excel(writer, "Прогнозы")

    return create_report_record(
        db=db,
        file_name=file_name,
        file_path=file_path,
        period_start=period_start,
        period_end=period_end,
    )


def make_temp_png_path():
    file_descriptor, file_path = tempfile.mkstemp(suffix=".png")
    os.close(file_descriptor)
    return file_path


def create_price_chart(product: Product, prices_df: pd.DataFrame):
    if prices_df.empty:
        return None

    chart_df = prices_df.copy()

    chart_df["Дата"] = pd.to_datetime(
        chart_df["Дата"],
        format="%d.%m.%Y %H:%M",
        errors="coerce"
    )

    chart_df["Цена"] = normalize_numeric_series(chart_df["Цена"])

    chart_df = chart_df.dropna(subset=["Дата", "Цена"]).sort_values("Дата")

    if chart_df.empty:
        return None

    fig = px.line(
        chart_df,
        x="Дата",
        y="Цена",
        markers=True,
        title=f"История цен: {product.name}",
        labels={
            "Дата": "Дата",
            "Цена": "Цена, руб."
        }
    )

    fig.update_layout(
        template="plotly_white",
        width=1200,
        height=600,
        xaxis_title="Дата",
        yaxis_title="Цена, руб.",
        margin=dict(l=80, r=40, t=80, b=80),
    )

    fig.update_xaxes(tickformat="%d.%m.%Y")
    fig.update_yaxes(tickformat=".0f")

    image_path = make_temp_png_path()
    fig.write_image(image_path)

    return image_path


def create_forecast_chart(
    product: Product,
    prices_df: pd.DataFrame,
    forecasts_df: pd.DataFrame
):
    if prices_df.empty or forecasts_df.empty:
        return None

    history_df = prices_df.copy()
    history_df["Дата"] = pd.to_datetime(
        history_df["Дата"],
        format="%d.%m.%Y %H:%M",
        errors="coerce"
    )
    history_df["Цена"] = normalize_numeric_series(history_df["Цена"])
    history_df = history_df.dropna(subset=["Дата", "Цена"]).sort_values("Дата")

    forecast_df = forecasts_df.copy()
    forecast_df["Дата прогноза"] = pd.to_datetime(
        forecast_df["Дата прогноза"],
        format="%d.%m.%Y",
        errors="coerce"
    )
    forecast_df["Прогнозируемая цена"] = normalize_numeric_series(
        forecast_df["Прогнозируемая цена"]
    )
    forecast_df = forecast_df.dropna(
        subset=["Дата прогноза", "Прогнозируемая цена"]
    ).sort_values("Дата прогноза")

    if history_df.empty or forecast_df.empty:
        return None

    fig = px.line(
        history_df,
        x="Дата",
        y="Цена",
        markers=True,
        title=f"Прогнозирование цены: {product.name}",
        labels={
            "Дата": "Дата",
            "Цена": "Цена, руб."
        }
    )

    fig.add_scatter(
        x=forecast_df["Дата прогноза"],
        y=forecast_df["Прогнозируемая цена"],
        mode="lines+markers",
        name="Прогноз",
        line={"dash": "dash"}
    )

    fig.update_layout(
        template="plotly_white",
        width=1200,
        height=600,
        xaxis_title="Дата",
        yaxis_title="Цена, руб.",
        margin=dict(l=80, r=40, t=80, b=80),
    )

    fig.update_xaxes(tickformat="%d.%m.%Y")
    fig.update_yaxes(tickformat=".0f")

    image_path = make_temp_png_path()
    fig.write_image(image_path)

    return image_path


def create_correlation_chart(product: Product, correlations_df: pd.DataFrame):
    if correlations_df.empty:
        return None

    required_columns = [
        "ID первой карточки",
        "ID второй карточки",
        "Коэффициент корреляции",
    ]

    for column in required_columns:
        if column not in correlations_df.columns:
            return None

    matrix_data = correlations_df[required_columns].copy()

    matrix_data["Коэффициент корреляции"] = normalize_numeric_series(
        matrix_data["Коэффициент корреляции"]
    )

    matrix_data = matrix_data.dropna()

    if matrix_data.empty:
        return None

    labels = sorted(
        set(matrix_data["ID первой карточки"].tolist())
        |
        set(matrix_data["ID второй карточки"].tolist())
    )

    matrix = pd.DataFrame(
        1.0,
        index=labels,
        columns=labels
    )

    for _, row in matrix_data.iterrows():
        first_id = row["ID первой карточки"]
        second_id = row["ID второй карточки"]
        value = row["Коэффициент корреляции"]

        matrix.loc[first_id, second_id] = value
        matrix.loc[second_id, first_id] = value

    fig = px.imshow(
        matrix,
        text_auto=True,
        title=f"Матрица корреляций: {product.name}",
        labels={
            "x": "Карточка товара",
            "y": "Карточка товара",
            "color": "Корреляция"
        },
        zmin=-1,
        zmax=1
    )

    fig.update_layout(
        template="plotly_white",
        width=1000,
        height=700,
        margin=dict(l=80, r=40, t=80, b=80),
    )

    image_path = make_temp_png_path()
    fig.write_image(image_path)

    return image_path


def prepare_pdf_table_data(dataframe: pd.DataFrame, max_rows: int = 20):
    if dataframe.empty:
        return [["Сообщение"], ["Данные отсутствуют"]]

    limited_df = dataframe.head(max_rows)

    return [limited_df.columns.tolist()] + limited_df.astype(str).values.tolist()


def add_table_to_pdf(elements, dataframe, font_name: str):
    table_data = prepare_pdf_table_data(dataframe)

    table = Table(table_data, repeatRows=1)

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )

    elements.append(table)


def add_image_to_pdf(elements, image_path: str):
    if image_path and os.path.exists(image_path):
        elements.append(Image(image_path, width=720, height=360))
        elements.append(Spacer(1, 18))


def add_missing_chart_text(elements, text: str, styles):
    elements.append(
        Paragraph(
            text,
            styles["BodyText"]
        )
    )
    elements.append(Spacer(1, 18))


def generate_pdf_report(db: Session, product_id: int):
    ensure_reports_directory()

    product = get_product(db, product_id)

    if product is None:
        raise ValueError("Товар не найден")

    font_name = register_pdf_font()
    product_name = safe_filename(product.name)
    created_at_text = get_current_datetime().strftime("%d-%m-%Y_%H-%M-%S")

    file_name = f"Отчёт_{product_name}_{created_at_text}.pdf"
    file_path = os.path.join(REPORTS_DIR, file_name)

    prices_df = get_prices_dataframe(db, product_id)
    anomalies_df = get_anomalies_dataframe(db, product_id)
    correlations_df = get_correlations_dataframe(db, product_id)
    forecasts_df = get_forecasts_dataframe(db, product_id)

    period_start, period_end = get_period(prices_df)

    temp_images = []

    try:
        price_chart = create_price_chart(product, prices_df)
        forecast_chart = create_forecast_chart(product, prices_df, forecasts_df)
        correlation_chart = create_correlation_chart(product, correlations_df)

        temp_images = [
            image_path
            for image_path in [price_chart, forecast_chart, correlation_chart]
            if image_path
        ]

        document = SimpleDocTemplate(
            file_path,
            pagesize=landscape(A4),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25,
        )

        styles = getSampleStyleSheet()
        styles["Title"].fontName = font_name
        styles["Heading2"].fontName = font_name
        styles["BodyText"].fontName = font_name

        elements = []

        elements.append(
            Paragraph(
                f"Отчёт по товару: {product.name}",
                styles["Title"]
            )
        )
        elements.append(Spacer(1, 12))

        elements.append(
            Paragraph(
                f"Дата формирования отчёта: {get_current_datetime().strftime('%d.%m.%Y %H:%M')}",
                styles["BodyText"]
            )
        )
        elements.append(Spacer(1, 18))

        chart_sections = [
            (
                "График истории цен",
                price_chart,
                "Нет данных для построения графика истории цен"
            ),
            (
                "График прогнозирования",
                forecast_chart,
                "Нет данных для построения графика прогнозирования"
            ),
            (
                "Матрица корреляций",
                correlation_chart,
                "Нет данных для построения матрицы корреляций"
            ),
        ]

        for index, (section_title, image_path, empty_text) in enumerate(chart_sections):
            if index > 0:
                elements.append(PageBreak())

            elements.append(Paragraph(section_title, styles["Heading2"]))
            elements.append(Spacer(1, 12))

            if image_path:
                add_image_to_pdf(elements, image_path)
            else:
                add_missing_chart_text(elements, empty_text, styles)

        table_sections = [
            ("История цен", prices_df),
            ("Выявленные аномалии", anomalies_df),
            ("Корреляционный анализ", correlations_df),
            ("Прогнозирование", forecasts_df),
        ]

        for section_title, dataframe in table_sections:
            elements.append(PageBreak())
            elements.append(Paragraph(section_title, styles["Heading2"]))
            elements.append(Spacer(1, 8))
            add_table_to_pdf(elements, dataframe, font_name)
            elements.append(Spacer(1, 18))

        document.build(elements)

        return create_report_record(
            db=db,
            file_name=file_name,
            file_path=file_path,
            period_start=period_start,
            period_end=period_end,
        )

    finally:
        for image_path in temp_images:
            if os.path.exists(image_path):
                os.remove(image_path)


def delete_report(db: Session, report_id: int):
    report = db.query(Report).filter(Report.id == report_id).first()

    if report is None:
        raise ValueError("Отчёт не найден")

    if os.path.exists(report.file_path):
        os.remove(report.file_path)

    db.delete(report)
    db.commit()